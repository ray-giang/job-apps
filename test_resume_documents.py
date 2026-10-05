import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from resume_documents import edit_copy, find_jobs


class ResumeDocumentTests(unittest.TestCase):
    def test_review_highlights_changes_without_changing_page_geometry(self):
        template = Path('/Users/air/Documents/Career/Resume/Raymond Giang Analytics Manager ATS.docx')
        master = json.loads((Path(__file__).parent / 'resume_master.json').read_text())
        from resume_tailor import select_resume
        job = {'company': 'Example', 'title': 'Staff Data Engineer', 'description': 'SQL dbt Snowflake data models'}
        selection = select_resume(master, job, max_bullets=18)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'review.docx'
            edit_copy(template, output, selection, review=True)
            source, result = Document(template), Document(output)
            self.assertEqual(len(source.sections), len(result.sections))
            self.assertEqual(source.sections[0].page_width, result.sections[0].page_width)
            self.assertEqual(source.sections[0].left_margin, result.sections[0].left_margin)
            def shaded(run):
                element = run._r.get_or_add_rPr().find(qn('w:shd'))
                return element is not None and element.get(qn('w:fill')) == 'FFF200'
            self.assertTrue(any(shaded(run) for paragraph in result.paragraphs for run in paragraph.runs))
            highlighted = sum(shaded(run) for paragraph in result.paragraphs for run in paragraph.runs)
            self.assertLessEqual(highlighted, 5)

    def test_find_jobs_supports_exact_urls_for_daily_new_jobs(self):
        with tempfile.TemporaryDirectory() as directory:
            jobs = Path(directory) / 'jobs.csv'
            jobs.write_text(
                'company,title,url,overall_confidence,compensation_confidence\n'
                'A,Role A,https://example.com/a,90,12\n'
                'B,Role B,https://example.com/b,95,10\n',
                encoding='utf-8',
            )
            selected = find_jobs(jobs, 3, ['https://example.com/a'])
            self.assertEqual([row['company'] for row in selected], ['A'])
            with self.assertRaises(ValueError):
                find_jobs(jobs, 3, ['https://example.com/missing'])


if __name__ == '__main__':
    unittest.main()

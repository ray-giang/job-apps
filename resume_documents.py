#!/usr/bin/env python3
"""Create clean and review resume DOCX copies without changing template formatting."""

import argparse
import csv
import json
import re
import shutil
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from resume_tailor import select_resume

ROOT = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = Path("/Users/air/Documents/Career/Resume/Raymond Giang Analytics Manager ATS.docx")

SUMMARY = {
    "management": "Analytics leader with 9+ years of experience supporting Product, Marketing, and Operations from problem definition through executive-ready insights. Hands-on manager with experience leading and developing analysts, defining quarterly roadmaps, establishing KPI governance, and building modern analytics systems with SQL, Snowflake, dbt, Airflow, GitHub, and Looker. Proven record delivering experimentation, attribution, monetization, and forecasting work that drives growth and operating efficiency.",
    "marketing": "Marketing analytics and measurement leader with 9+ years of experience supporting Product, Marketing, and Operations from problem definition through executive-ready insights. Hands-on technical partner with experience in experimentation, multi-touch attribution, first-party data, segmentation, forecasting, and modern analytics systems using SQL, Snowflake, dbt, Airflow, Python, and Looker. Proven record translating analytical work into growth, monetization, and operating decisions.",
    "product": "Product analytics leader with 9+ years of experience guiding experimentation, monetization, retention, and measurement decisions across Product, Marketing, and Operations. Hands-on technical partner with experience developing analysts, defining analytics roadmaps, establishing KPI governance, and building modern analytics systems using SQL, Snowflake, dbt, Airflow, Python, and Looker. Proven record turning product behaviour and experimental evidence into measurable growth.",
    "data_science": "Data science and analytics leader with 9+ years of experience applying experimentation, segmentation, forecasting, attribution, and statistical analysis to Product, Marketing, and Operations decisions. Hands-on technical partner with experience in Python, SQL, Snowflake, dbt, Airflow, and Looker, plus leadership experience developing analysts and analytics roadmaps. Proven record translating models and experiments into measurable growth and operating outcomes.",
    "analytics_engineering": "Analytics engineering and data leader with 9+ years of experience building trusted data models, pipelines, reporting systems, and analytical products for Product, Marketing, and Operations. Hands-on technical expertise includes SQL, Snowflake, dbt, Airflow, Python, GitHub, Looker, Fivetran, and data modelling, alongside experience developing analysts and analytics roadmaps. Proven record improving scalability, data quality, attribution, forecasting, and business decision-making.",
    "general": "Analytics and data professional with 9+ years of experience supporting Product, Marketing, and Operations from problem definition through executive-ready insights. Hands-on technical and people leader with experience in modern analytics systems, experimentation, attribution, forecasting, KPI governance, and team development. Proven record delivering measurable growth and operating improvements."
}


def shade_change(run):
    """Apply review shading without Word's highlight markup, which breaks LO layout."""
    properties = run._r.get_or_add_rPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), "FFF200")

TRACK_BULLET_REPLACEMENTS = {
    "analytics_engineering": {21: "opencare_forecast"},
    "data_science": {22: "opencare_forecast"},
}


def replace_text(paragraph, text, highlight=False):
    original = paragraph.text
    runs = paragraph.runs
    if not runs:
        run = paragraph.add_run(text)
    else:
        run = runs[0]
        run.text = text
        for extra in runs[1:]:
            extra.text = ""
    if highlight and original.strip() != text.strip():
        shade_change(run)


def replace_headline(paragraph, headline, highlight=False):
    runs = [run for run in paragraph.runs if run.text.strip()]
    if not runs:
        raise ValueError("Template title paragraph has no editable runs")
    target = runs[-1]
    changed = target.text.strip() != headline.strip()
    target.text = headline
    if highlight and changed:
        shade_change(target)


def edit_copy(template, output, selection, review=False):
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, output)
    doc = Document(output)
    paragraphs = doc.paragraphs
    if len(paragraphs) < 36:
        raise ValueError("Resume template structure changed; expected at least 36 paragraphs")
    replace_headline(paragraphs[0], selection["headline"], review)
    replace_text(paragraphs[2], SUMMARY[selection["track"]], review)
    # Keep this centered multi-line paragraph unshaded: LibreOffice can corrupt
    # the surrounding floating layout when run shading spans a wrapped line.
    replace_text(paragraphs[4], "| " + " | ".join(selection["skills"]) + " |", False)
    evidence = selection["evidence_library"]
    for slot, evidence_id in TRACK_BULLET_REPLACEMENTS.get(selection["track"], {}).items():
        if evidence_id not in evidence:
            raise ValueError(f"Missing resume evidence {evidence_id}")
        replace_text(paragraphs[slot], evidence[evidence_id], review)
    doc.save(output)


def safe_name(value):
    return re.sub(r"[^A-Za-z0-9&+.-]+", " ", value).strip()


def find_jobs(path, count, job_urls=None):
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    ranked = sorted(rows, key=lambda row: (int(row["overall_confidence"]), int(row["compensation_confidence"])), reverse=True)
    if not job_urls:
        return ranked[:count]
    requested = set(job_urls)
    selected = [row for row in ranked if row.get("url") in requested]
    missing = requested - {row.get("url") for row in selected}
    if missing:
        raise ValueError(f"Job URLs not found in jobs.csv: {sorted(missing)}")
    return selected


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--jobs", type=Path, default=ROOT / "jobs.csv")
    parser.add_argument("--master", type=Path, default=ROOT / "resume_master.json")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--clean-dir", type=Path, required=True)
    parser.add_argument("--top", type=int, default=3)
    parser.add_argument("--job-url", action="append", help="Generate only this exact jobs.csv URL; repeat for multiple jobs")
    args = parser.parse_args()
    master = json.loads(args.master.read_text(encoding="utf-8"))
    manifest = []
    jobs = find_jobs(args.jobs, args.top, args.job_url)
    company_counts = Counter(safe_name(job["company"]) for job in jobs)
    for job in jobs:
        selection = select_resume(master, job, max_bullets=18)
        company = safe_name(job["company"])
        role_suffix = f" {safe_name(job['title'])}" if company_counts[company] > 1 else ""
        stem = f"Raymond Giang {company}{role_suffix} Resume"
        review_path = args.output_dir / f"{stem}.docx"
        clean_path = args.clean_dir / f"{stem}.docx"
        edit_copy(args.template, clean_path, selection, review=False)
        edit_copy(args.template, review_path, selection, review=True)
        manifest.append({"company": job["company"], "title": job["title"], "score": int(job["overall_confidence"]),
                         "track": selection["track"], "clean_docx": str(clean_path), "review_docx": str(review_path),
                         "pdf": str(args.output_dir / f"{stem}.pdf"), "url": job["url"], "evidence_ids": selection["evidence_ids"]})
    manifest_path = args.clean_dir / "resume_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(manifest_path)


if __name__ == "__main__":
    main()
